import io
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi.responses import StreamingResponse


def generate_envelope_dl(transaction) -> StreamingResponse:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(220)
    section.page_height = Mm(110)
    section.left_margin = Mm(10)
    section.right_margin = Mm(10)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(10)

    plantation = transaction.plot.plantation

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_title.add_run(plantation.organization.name)
    run.font.size = Pt(20)
    run.font.bold = True

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    owner_names = ", ".join(s.owner.username for s in transaction.owner_shares)
    p_info.add_run(
        f"วันที่: {transaction.date.strftime('%d/%m/%Y')}  |  "
        f"สวน: {plantation.name}  |  แปลง: {transaction.plot.name}\n"
        f"เจ้าของ: {owner_names}  |  คนกรีด: {transaction.tapper.name}"
    ).font.size = Pt(12)

    p_amount = doc.add_paragraph()
    p_amount.alignment = WD_ALIGN_PARAGRAPH.CENTER
    owners_total = sum(s.total_amount for s in transaction.owner_shares)
    run_amount = p_amount.add_run(f"NET เจ้าของ: {round(owners_total):,} บาท")
    run_amount.font.size = Pt(36)
    run_amount.font.bold = True

    p_tapper = doc.add_paragraph()
    p_tapper.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tapper.add_run(
        f"คนกรีด: {round(transaction.tapper_share):,} บาท"
    ).font.size = Pt(18)

    return _stream_doc(doc, f"envelope_{transaction.id}.docx")


def generate_a4_summary(transaction) -> StreamingResponse:
    doc = Document()
    plantation = transaction.plot.plantation

    doc.add_heading(f"{plantation.organization.name} — ใบสรุปรายการ", 0)

    # Header info
    p = doc.add_paragraph()
    p.add_run(f"รายการที่: #{transaction.id}\n").bold = True
    p.add_run(f"วันที่: {transaction.date.strftime('%d/%m/%Y %H:%M')}\n")
    p.add_run(f"สวน: {plantation.name}  |  แปลง: {transaction.plot.name}\n")
    p.add_run(f"คนกรีด: {transaction.tapper.name}\n")
    p.add_run(f"ผู้บันทึก: {transaction.buyer.username}")

    # Rubber calculation table
    doc.add_heading("รายละเอียดยาง", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "รายการ"
    table.rows[0].cells[1].text = "จำนวน (บาท)"

    rows = [
        ("น้ำหนัก (กก.)", f"{transaction.rubber_weight:.2f}"),
        ("ราคา/กก.", f"{transaction.rubber_price_per_unit:.2f}"),
        ("รวมยาง (Gross)", f"{transaction.gross_rubber_amount:.2f}"),
        ("ส่วนคนกรีด 50%", f"{transaction.tapper_share:.2f}"),
        ("ส่วนเจ้าของสวน 50%", f"{transaction.owner_rubber_pool:.2f}"),
        (f"หักค่าทีมงาน {plantation.backteam_percentage}%", f"-{transaction.backteam_amount:.2f}"),
        ("ยางสุทธิของเจ้าของ", f"{transaction.owner_rubber_net:.2f}"),
    ]
    if transaction.wood_income > 0:
        rows.append(("ค่าไม้ (100% เจ้าของ)", f"+{transaction.wood_income:.2f}"))

    for label, value in rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value

    # Owner breakdown
    doc.add_heading("แบ่งให้เจ้าของสวน", level=2)
    table2 = doc.add_table(rows=1, cols=4)
    table2.style = "Table Grid"
    hdr = table2.rows[0].cells
    hdr[0].text = "เจ้าของ"
    hdr[1].text = "สิทธิ์ %"
    hdr[2].text = "ยาง (บาท)"
    hdr[3].text = "รวม (บาท)"
    for s in transaction.owner_shares:
        row = table2.add_row().cells
        row[0].text = s.owner.username
        row[1].text = f"{s.share_percentage:.0f}%"
        row[2].text = f"{s.rubber_amount:.2f}"
        row[3].text = f"{round(s.total_amount):,}"

    # BackTeam breakdown
    doc.add_heading("แบ่งให้ทีมงาน", level=2)
    table3 = doc.add_table(rows=1, cols=3)
    table3.style = "Table Grid"
    hdr3 = table3.rows[0].cells
    hdr3[0].text = "สมาชิก"
    hdr3[1].text = "สัดส่วน %"
    hdr3[2].text = "จำนวน (บาท)"
    for s in transaction.backteam_shares:
        row = table3.add_row().cells
        row[0].text = s.member.username
        row[1].text = f"{s.share_percentage:.0f}%"
        row[2].text = f"{s.amount:.2f}"

    # Grand total
    doc.add_heading("สรุป (ปัดเศษเพื่อจ่ายเงินสด)", level=2)
    p_total = doc.add_paragraph()
    p_total.add_run(f"คนกรีดรับ: {round(transaction.tapper_share):,} บาท\n").bold = True
    owners_total = sum(s.total_amount for s in transaction.owner_shares)
    p_total.add_run(f"เจ้าของสวนรวม: {round(owners_total):,} บาท\n").bold = True
    backteam_total = sum(s.amount for s in transaction.backteam_shares)
    p_total.add_run(f"ทีมงานรวม: {backteam_total:.2f} บาท").bold = True

    return _stream_doc(doc, f"summary_a4_{transaction.id}.docx")


def _stream_doc(doc: Document, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
